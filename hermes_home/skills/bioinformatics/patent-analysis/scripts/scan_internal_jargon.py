#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""交付件「内部管理信息残留」扫描探针（纯本地，依赖 python-docx）

触发场景：交底书 / 附件 / 任何对外交付的 docx 生成完成后，
          自查里面有没有残留**项目组内部才有意义**的信息：
          版本行、版本别名（"正文瘦身版"）、数据口径行、内部战略代号、
          内部文件名（claims.md）、内部交付路径（交付目录 / run_all.sh）。

用法：
    python scan_internal_jargon.py --docs 技术交底书_v16.docx 附件1_….docx 附件2_….docx
    python scan_internal_jargon.py --docs a.docx --extra "项目代号X" --extra "内部批次"
    python scan_internal_jargon.py --docs a.docx --keys 战略 口径 --no-advisory

🔴 扫描范围必须含四处 —— 漏一处就漏一类：
    ① 正文段落  ② 表格单元格  ③ 页眉  ④ 页脚
   （实测标本：版本行被"正文瘦身"顺手搬进了**页脚**，只扫正文会报 0 命中然后被用户当场抓到）

判据：HARD 命中 = 0 才算通过。命中后**改生成器脚本**（删常量 / 加 extra_edits 映射），
      不要手改 docx —— 手改的下一轮会被生成器覆盖回去。

退出码：0 = 无 HARD 命中；1 = 有 HARD 命中（可直接用于验收脚本串联）。
"""
from __future__ import annotations

import argparse
import os
import sys

try:
    import docx  # python-docx
except ImportError:
    sys.stderr.write('缺少 python-docx。持久内核没有时改用系统 Python312 跑，不要为此重复装包。\n')
    raise SystemExit(2)

# ── 必须清除的内部管理信息（按类别分组，便于汇报时说明"清的是哪一类"） ──
HARD_KEYS = [
    # 版本行 / 版本别名 / 数据口径
    '版本 v', '版本v', '版本：', '版本:', '瘦身', '口径：', '口径:',
    # 内部战略 / 方案代号
    '战略 ', '战略A', '战略B', '战略C', '方案 A+', '草案 v', '草案v',
    # 内部文件名
    'claims.md', 'disclosure.md', 'abstract.md', 'README.md', 'A22.3', 'run_all.sh',
    # 内部交付语 / 内部自称
    '交付目录', '非交底书字段', '本会话', 'memomics', 'session', '自动生成',
]

# ── 考虑到但**允许保留**的词：只提示、不算失败（技术对照标签 / 法条援引） ──
ADVISORY_KEYS = ['基线①', '基线②', '基线③', 'A22.4', 'A25', 'A26.4', '支撑独立权利要求', '供参考']


def iter_text_units(doc):
    """产出 (区域标记, 文本)；区域标记形如 P12 / T0R1C2 / 页眉 / 页脚。"""
    for i, p in enumerate(doc.paragraphs):
        yield f'P{i}', p.text
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                yield f'T{ti}R{ri}C{ci}', cell.text
    for si, sec in enumerate(doc.sections):
        for tag, part in (('页眉', sec.header), ('页脚', sec.footer)):
            label = tag if len(doc.sections) == 1 else f'{tag}{si}'
            for p in part.paragraphs:
                yield label, p.text
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield label, cell.text


def scan(path, keys, advisory):
    d = docx.Document(path)
    units = list(iter_text_units(d))
    hits, adv = [], []
    for key in keys:
        for where, text in units:
            if key in text:
                hits.append((key, where, text.strip().replace('\n', ' ')[:160]))
    for key in advisory:
        for where, text in units:
            if key in text:
                adv.append((key, where, text.strip().replace('\n', ' ')[:120]))
    return d, hits, adv


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--docs', nargs='+', required=True, help='待扫描的 docx 路径（可多个）')
    ap.add_argument('--keys', nargs='*', default=HARD_KEYS, help='覆盖 HARD 关键词表')
    ap.add_argument('--extra', action='append', default=[], help='追加关键词（可重复）')
    ap.add_argument('--no-advisory', action='store_true', help='不打印"考虑到但允许保留"的提示段')
    args = ap.parse_args()

    keys = list(args.keys) + list(args.extra)
    advisory = [] if args.no_advisory else ADVISORY_KEYS
    total_hard = 0

    for path in args.docs:
        if not os.path.exists(path):
            print(f'[跳过] 文件不存在：{path}\n')
            continue
        d, hits, adv = scan(path, keys, advisory)
        total_hard += len(hits)
        print('=' * 78)
        print(f'{os.path.basename(path)}  {os.path.getsize(path):,} B  '
              f'段{len(d.paragraphs)} 表{len(d.tables)} 图{len(d.inline_shapes)} 节{len(d.sections)}')
        if hits:
            print(f'  🔴 内部管理信息命中 {len(hits)} 处（必须清除）：')
            for key, where, text in hits:
                print(f'     [{key}] {where}: {text}')
        else:
            print('  ✅ 内部管理信息命中 0 处（段/表/页眉/页脚四处已扫）')
        if adv:
            print(f'  🟡 以下词命中 {len(adv)} 处，属"技术对照标签 / 法条援引"，**默认保留**，供你确认：')
            for key, where, text in adv[:12]:
                print(f'     [{key}] {where}: {text}')
            if len(adv) > 12:
                print(f'     … 另有 {len(adv) - 12} 处')
        print()

    print('=' * 78)
    if total_hard:
        print(f'结论：FAIL —— 共 {total_hard} 处内部管理信息残留。'
              f'修法=改生成器（删常量 + 加 extra_edits 映射）后重跑，禁止手改 docx。')
    else:
        print('结论：PASS —— 全部件内部管理信息 0 命中（版本行/口径/战略代号/内部文件名/交付路径）。')
    return 1 if total_hard else 0


if __name__ == '__main__':
    raise SystemExit(main())
