# -*- coding: utf-8 -*-
"""
交付目录 ↔ 交底书 docx 对应性核查（第 2、3 刀）:
  ① 从 docx 抽出内嵌图，并按 document.xml 的 r:embed 顺序校验"图 ↔ 图注"配对
  ② 复核文档"复现脚本清单"里列出的每个脚本文件名是否真实存在于交付目录

用法:
    python check_docx_figure_pairing.py <docx路径> <交付目录> [抽取输出目录]

要点:
  - 图 ↔ 图注配对**不能靠文件名猜**，必须读 document.xml 里图片所在段落的
    前一段/后一段文字（本类文档实测为"图在前、图注在后"）。
  - 第 ② 刀是最常漂移的一项：文档写的脚本名往往早已被改名/升版本，
    照表跑会直接报文件不存在。（实测：07_age_shuffle.R 实为
    07_age_shuffle_permutation.R；08_gen_..._v10.py 实为 ..._v11.py）
"""
import os
import re
import sys
import zipfile


def extract_and_pair(docx, outdir):
    os.makedirs(outdir, exist_ok=True)
    with zipfile.ZipFile(docx) as z:
        media = [n for n in z.namelist() if n.startswith('word/media/')]
        z.extractall(outdir)
        rels = z.read('word/_rels/document.xml.rels').decode('utf-8', 'ignore')
        doc = z.read('word/document.xml').decode('utf-8', 'ignore')

    rid2img = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="media/([^"]+)"', rels))
    n_img = len(media)
    print('=' * 62)
    print('步骤 ① · 内嵌图与图注配对校验')
    print('=' * 62)
    print(f'  word/media/ 内嵌图数: {n_img}')

    paras = re.findall(r'<w:p[ >].*?</w:p>', doc, re.S)

    def txt(p):
        return ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', p))

    order = [rid2img[r] for r in re.findall(r'r:embed="(rId\d+)"', doc) if r in rid2img]
    print(f'  document.xml 引用顺序: {order}')
    for i, p in enumerate(paras):
        if 'r:embed' not in p:
            continue
        img = rid2img.get(re.findall(r'r:embed="(rId\d+)"', p)[0], '?')
        before = txt(paras[i - 1])[:60] if i > 0 else ''
        after = txt(paras[i + 1])[:60] if i + 1 < len(paras) else ''
        print(f'\n  [{img}]')
        print(f'      前一段: {before!r}')
        print(f'      后一段: {after!r}')
    print('\n  ⇒ 判读：图注在图的"后一段"则配对正确；数量不符或错位即为不一致。')
    return outdir


SCRIPT_RE = re.compile(r'(?:scripts/)?([0-9A-Za-z_\u4e00-\u9fff]+\.(?:py|R|sh))')


def check_script_inventory(docx, delivery):
    print()
    print('=' * 62)
    print('步骤 ② · 文档"复现脚本清单" ↔ 交付目录实际文件')
    print('=' * 62)
    try:
        import docx as _docx  # noqa: F401
        have_docx = True
    except ImportError:
        have_docx = False

    if have_docx:
        d = _docx.Document(docx)
        text = '\n'.join(p.text for p in d.paragraphs)
        for t in d.tables:
            for row in t.rows:
                text += '\n' + '\t'.join(c.text for c in row.cells)
    else:
        with zipfile.ZipFile(docx) as z:
            text = z.read('word/document.xml').decode('utf-8', 'ignore')
        text = ' '.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', text))

    cited = sorted(set(SCRIPT_RE.findall(text)))
    actual = set()
    for root, _, files in os.walk(delivery):
        for f in files:
            actual.add(f)

    missing, ok = [], []
    for c in cited:
        base = os.path.basename(c)
        (ok if base in actual else missing).append(base)

    print(f'  文档提到的脚本: {len(cited)} 个')
    for c in cited:
        base = os.path.basename(c)
        print(f'    {"✓" if base in actual else "✗ 不存在"}  {base}')
    if missing:
        print(f'\n  ⚠️ {len(missing)} 个脚本在交付目录中不存在 —— 照表跑会报文件不存在:')
        for m in missing:
            print(f'      - {m}')
    print(f'\n  ⇒ 另需人工核对: 生成器版本号是否与当前 docx 一致；'
          f'文档声称"出图脚本产出附图 1–N"时，读该脚本的 savefig 调用数它实际产出几张。')
    return 0


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 2
    docx, delivery = sys.argv[1], sys.argv[2]
    outdir = sys.argv[3] if len(sys.argv) > 3 else os.path.join(os.getcwd(), '_docx_media_check')
    extract_and_pair(docx, outdir)
    check_script_inventory(docx, delivery)
    print(f'\n内嵌图已抽取到: {outdir}\\word\\media\\')
    print('⚠️ Windows: 视觉读图工具需要 Windows 绝对路径，MSYS 的 /tmp/... 会报"图片不存在"——'
          '先 cp 到项目目录再读。')
    return 0


if __name__ == '__main__':
    sys.exit(main())
