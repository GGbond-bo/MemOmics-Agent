# -*- coding: utf-8 -*-
"""
中文 PPTX 构建脚手架（浅色学术风 · python-pptx 原生形状 · 全部可编辑）
================================================================
用法：拷成 results/<sid>/scripts/build_xxx_ppt.py，改 CONTENT 区即可。
要点：
  · 中文内容一律用 Microsoft YaHei（Arial 无 CJK 字形 → 豆腐块）
  · 图内说明文字若走 matplotlib 必须英文；中文示意图用本文件的 rect/hline 原生形状
  · 文案里的中文引号用「」，不要在双引号字符串里嵌 ASCII 双引号（会 SyntaxError）
  · 输出到 results/<sid>/reports/，不要写 /mnt/results/ 或 /tmp/results-staging/
依赖：python-pptx 必须装进持久内核的解释器（先 print(sys.executable) 确认）
      terminal: ./.venv/Scripts/python.exe -m pip install python-pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------- 输出 ----------------
OUT_DIR = r"E:/MemOmics-Agent/results/<session_id>/reports"
OUT = os.path.join(OUT_DIR, "deck.pptx")
os.makedirs(OUT_DIR, exist_ok=True)

# ---------------- 浅色学术风配色 ----------------
INK     = RGBColor(0x1A, 0x1A, 0x1A)   # 主文字（深墨）
SUB     = RGBColor(0x5A, 0x64, 0x72)   # 次级文字
LINE    = RGBColor(0xD8, 0xDE, 0xE6)   # 分隔线
PANEL   = RGBColor(0xF5, 0xF7, 0xFA)   # 浅面板
PANEL2  = RGBColor(0xEE, 0xF3, 0xF9)   # 浅蓝面板
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
BLUE    = RGBColor(0x2E, 0x6F, 0xB7)
TEAL    = RGBColor(0x1F, 0x8A, 0x70)
AMBER   = RGBColor(0xC7, 0x7A, 0x2B)
PURPLE  = RGBColor(0x6B, 0x4F, 0xA8)
CRIMSON = RGBColor(0xB0, 0x3A, 0x48)   # 主强调色
GREEN   = RGBColor(0x2F, 0x7D, 0x4F)
RED     = RGBColor(0xA8, 0x32, 0x32)

FONT = "Microsoft YaHei"   # ← 中文 deck 必改；纯英文 deck 才用 "Arial"

prs = Presentation()
prs.slide_width = Inches(13.333)      # 16:9
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
MARGIN = 0.55
CW = 12.23                            # 内容宽 = 13.333 - 2*0.55


# ================= 通用辅助 =================
def tb(slide, l, t, w, h, text, size=14, color=INK, bold=False,
       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0, wrap=True):
    """文本框。text 可含 \\n，每行一段。"""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = FONT
    return box


def rect(slide, l, t, w, h, fill=None, line=None,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=0.75, radius=0.08):
    """形状。fill=None → 透明；line=None → 无边框。中文示意图就靠它画。"""
    s = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    s.shadow.inherit = False          # 浅色学术风：一律去阴影
    return s


def hline(slide, l, t, w, color=LINE, width=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Pt(width))
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background(); s.shadow.inherit = False
    return s


def header(slide, num, title, sub=""):
    """统一页眉：左侧色块 + 标题 + 副题 + 分隔线 + 右上页码"""
    rect(slide, MARGIN, 0.42, 0.055, 0.52, fill=CRIMSON, shape=MSO_SHAPE.RECTANGLE)
    tb(slide, MARGIN + 0.23, 0.40, 9.6, 0.42, title, size=21, bold=True)
    if sub:
        tb(slide, MARGIN + 0.23, 0.86, 11.9, 0.28, sub, size=11.5, color=SUB)
    tb(slide, 12.30, 0.44, 0.5, 0.3, num, size=11, color=SUB, align=PP_ALIGN.RIGHT)
    hline(slide, MARGIN, 1.24, CW)


def footer(slide, text="", right="浅色学术风"):
    hline(slide, MARGIN, 7.00, CW)
    tb(slide, MARGIN, 7.08, 9.0, 0.24, text, size=9, color=SUB)
    tb(slide, 11.0, 7.08, 1.78, 0.24, right, size=9, color=SUB, align=PP_ALIGN.RIGHT)


def set_cell(cell, text, size=10.5, bold=False, color=INK, fill=None,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.06):
    cell.margin_left = cell.margin_right = Inches(margin)
    cell.margin_top = cell.margin_bottom = Inches(0.035)
    cell.vertical_anchor = anchor
    if fill is not None:
        cell.fill.solid(); cell.fill.fore_color.rgb = fill
    else:
        cell.fill.background()
    tf = cell.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color; r.font.name = FONT


def table(slide, l, t, w, rows, col_w, hdr_fill=PANEL2, row_h=0.42, hdr_h=0.40,
          size=10.5, hdr_size=10.5, body_fills=None):
    """rows[0] 为表头。body_fills={行序号: RGBColor} 用于高亮某行。"""
    nrow, ncol = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nrow, ncol, Inches(l), Inches(t), Inches(w),
                                Inches(hdr_h + row_h * (nrow - 1)))
    tx = gt.table
    for j, cwv in enumerate(col_w):
        tx.columns[j].width = Inches(cwv)
    tx.rows[0].height = Inches(hdr_h)
    for i in range(1, nrow):
        tx.rows[i].height = Inches(row_h)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            if i == 0:
                set_cell(tx.cell(i, j), val, size=hdr_size, bold=True, fill=hdr_fill)
            else:
                f = PANEL if (i % 2 == 0) else WHITE
                if body_fills and body_fills.get(i) is not None:
                    f = body_fills[i]
                set_cell(tx.cell(i, j), val, size=size, fill=f)
    return tx


def flow_row(slide, steps, y, x0=MARGIN, total_w=CW, bh=1.0,
             gap=0.20, box_fill=PANEL):
    """一行等宽流程框 + 箭头。steps=[(名称, 副文字, 颜色), ...]"""
    n = len(steps)
    bw = (total_w - gap * (n - 1)) / n
    x = x0
    for i, (name, sub, col) in enumerate(steps):
        rect(slide, x, y, bw, bh, fill=box_fill, line=col, line_w=1.15, radius=0.09)
        tb(slide, x + 0.10, y + 0.14, bw - 0.20, 0.30, name,
           size=11.5, bold=True, color=col, align=PP_ALIGN.CENTER)
        tb(slide, x + 0.08, y + 0.52, bw - 0.16, 0.42, sub,
           size=9.5, color=SUB, align=PP_ALIGN.CENTER, spacing=1.05)
        if i < n - 1:
            rect(slide, x + bw + 0.02, y + bh / 2 - 0.12, gap - 0.04, 0.24,
                 fill=LINE, shape=MSO_SHAPE.RIGHT_ARROW)
        x += bw + gap


def banner(slide, y, text, w=CW, fill=PANEL2, color=INK, size=10.5):
    """闭环/回流说明横幅（超宽圆角，视觉上当回路用）"""
    rect(slide, MARGIN, y, w, 0.40, fill=fill, radius=0.30)
    tb(slide, MARGIN, y + 0.06, w, 0.28, text, size=size, color=color,
       align=PP_ALIGN.CENTER)


# ================= CONTENT（改这里） =================
# --- P1 封面 ---
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 13.333, 7.5, fill=WHITE, shape=MSO_SHAPE.RECTANGLE)
rect(s, 0, 0, 0.22, 7.5, fill=CRIMSON, shape=MSO_SHAPE.RECTANGLE)
tb(s, 1.25, 1.12, 11.3, 0.75, "主标题", size=40, bold=True)
tb(s, 1.25, 1.92, 11.3, 0.5, "副标题", size=21, color=SUB)
hline(s, 1.25, 3.06, 11.0)
banner(s, 4.62, "一句话总纲（左侧红字强调版可另起一行 tb）", w=11.0, fill=PANEL2, size=14)

# --- P2 内容页 ---
s = prs.slides.add_slide(BLANK)
header(s, "02", "页标题", "副标题")
table(s, MARGIN, 1.46, CW,
      [["列1", "列2"], ["a", "b"]], col_w=[6.0, 6.23], row_h=0.5)
footer(s)

# --- P3 一行流程 ---
s = prs.slides.add_slide(BLANK)
header(s, "03", "流程图页", "原生形状 → 中文正常 + 可编辑")
flow_row(s, [("步骤一", "说明", BLUE), ("步骤二", "说明", TEAL),
             ("步骤三", "说明", RED)], y=1.52)
banner(s, 2.80, "↑ 回流说明")
footer(s)

# ================= 保存 + 结构校验 =================
prs.save(OUT)
size = os.path.getsize(OUT)
chk = Presentation(OUT)
n = len(chk.slides)
assert size > 10_000, f"文件仅 {size} 字节，疑似空稿"
assert n >= 2, f"仅 {n} 页，疑似缺内容"
for i, sl in enumerate(chk.slides, 1):
    assert len(sl.shapes) >= 1, f"第 {i} 页无形状，疑似空白"
bad = []
for i, sl in enumerate(chk.slides, 1):
    for sh in sl.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if run.font.name and run.font.name != FONT:
                    bad.append((i, run.font.name))
# 注意：这里只校验"字体是否为本 deck 的 FONT"，不再硬断言 Arial
print(f"[OK] {OUT} | {n} 页 | {size/1024:.1f} KB | 字体异常 {len(bad)} 处")
if bad:
    print("  ⚠ 字体不一致（中文可能豆腐块）:", bad[:5])