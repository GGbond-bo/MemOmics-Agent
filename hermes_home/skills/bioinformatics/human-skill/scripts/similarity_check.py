#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""查重自检（自建库 n-gram 口径）

用法:
    python similarity_check.py <论文.md> [--ref 目录或文件 ...] [--n 15] [--top 20] [--json out.json]

两项指标:
  1. 内部重复率 —— 同一 n-gram 在文内出现 ≥2 次的字符覆盖占比
  2. 源库重合率 —— 与 --ref 指定语料共有的 n-gram 字符覆盖占比（学位论文 vs 专利交底书等）

⚠ 口径声明：本工具是「自建库 n-gram 自检」，不是知网/CNKI。
   结果只能用于内部自查与改写定位，禁止当作官方查重率汇报。
"""
import os
import re
import sys
import json
import glob
import argparse
from collections import defaultdict


def read_text(path):
    """读 md/txt/docx 为纯文本行列表。"""
    if path.lower().endswith(".docx"):
        try:
            from docx import Document
        except ImportError:
            sys.exit("需要 python-docx 读 .docx（用装了 python-docx 的解释器，或改用 Markdown 母本）")
        d = Document(path)
        out = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                out.append(" ".join(c.text for c in row.cells))
        return out
    with open(path, encoding="utf-8", errors="ignore") as f:
        return f.read().splitlines()


def normalize(lines):
    """剥离 Markdown 标记/图片/代码块，全角半角统一，去空白。"""
    out = []
    in_code = False
    for L in lines:
        s = L.strip()
        if s.startswith("```"):
            in_code = not in_code
            continue
        if in_code or s.startswith("!["):
            continue
        s = re.sub(r"^\s*\|", "", s)
        s = re.sub(r"\|", " ", s)
        s = re.sub(r"[*_`#>\[\]]", "", s)          # markdown 标记
        s = re.sub(r"\$[^$]*\$", "", s)            # 行内公式
        s = s.translate(str.maketrans(
            "０１２３４５６７８９ＡＢＣＤＥＦＧＨＩＪＫＬＭＮＯＰＱＲＳＴＵＶＷＸＹＺ"
            "ａｂｃｄｅｆｇｈｉｊｋｌｍｎｏｐｑｒｓｔｕｖｗｘｙｚ（）：，。；",
            "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"
            "abcdefghijklmnopqrstuvwxyz():,.;"))
        s = re.sub(r"\s", "", s)
        if s:
            out.append(s)
    return out


def cjk_ratio(s):
    """片段中汉字占比（用于剔除英文/DOI/纯符号 n-gram，避免虚假重合）。"""
    return sum(1 for ch in s if "\u4e00" <= ch <= "\u9fff") / max(len(s), 1)


def ngrams(text, n, min_cjk=0.6):
    """取 n-gram；只保留汉字占比 ≥ min_cjk 的（默认 60%），过滤英文基因名/DOI/数字。"""
    out = []
    for i in range(0, max(len(text) - n + 1, 0)):
        g = text[i:i + n]
        if cjk_ratio(g) >= min_cjk:
            out.append(g)
    return out


def covered_chars(marks):
    """marks: 逐位置布尔列表 → 被标记的字符数。"""
    return sum(1 for m in marks if m)


def ensure_parent(path):
    import os as _os
    d = _os.path.dirname(_os.path.abspath(path))
    _os.makedirs(d, exist_ok=True)


def collect_dir(paths):
    files = []
    for p in paths:
        if os.path.isdir(p):
            for ext in ("*.md", "*.txt", "*.docx", "*.csv"):
                files += glob.glob(os.path.join(p, "**", ext), recursive=True)
        elif os.path.exists(p):
            files.append(p)
    return files


def longest_common_blocks(target, ref_grams, n, min_block, topn, min_cjk=0.6):
    """扫出 target 上与 ref 连续共有的片段。

    做法：逐位置判断 target[i:i+n] 是否在 ref 中，取连续 True 的段（run）。
    一段 run 长度为 L → 实际共有片段为 target[i : i+L+n-1]（L+n-1 个字符）。
    """
    marks = [target[i:i + n] in ref_grams for i in range(max(len(target) - n + 1, 0))]
    blocks, i, M = [], 0, len(marks)
    while i < M:
        if marks[i]:
            j = i
            while j < M and marks[j]:
                j += 1
            seg = target[i:j + n - 1]
            if len(seg) >= min_block and cjk_ratio(seg) >= min_cjk:
                blocks.append((len(seg), i, seg))
            i = j
        else:
            i += 1
    blocks.sort(reverse=True)
    return blocks[:topn]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("target")
    ap.add_argument("--ref", nargs="*", default=[])
    ap.add_argument("--n", type=int, default=15)
    ap.add_argument("--top", type=int, default=20)
    ap.add_argument("--json", default=None)
    a = ap.parse_args()

    body = "".join(normalize(read_text(a.target)))
    total = len(body)
    N = a.n

    # 位置标记法：重复 n-gram 覆盖的字符（不重复计数）
    pos_map = {}
    for i in range(max(total - N + 1, 0)):
        g = body[i:i + N]
        if cjk_ratio(g) >= 0.6:
            pos_map.setdefault(g, []).append(i)
    dup_mark = [False] * total
    intra = []
    for g, ps in pos_map.items():
        if len(ps) > 1:
            intra.append((len(ps), g))
            for i in ps:
                for k in range(i, min(i + N, total)):
                    dup_mark[k] = True
    dup_chars = covered_chars(dup_mark)
    rate = dup_chars / total * 100 if total else 0

    print("# 查重自检报告\n")
    print(f"**目标文件**：`{a.target}`")
    print(f"**归一化后正文字符数**：{total}")
    print(f"**n-gram 长度**：{N} 个汉字（仅统计汉字占比 ≥60% 的片段）\n")
    print("## 1. 内部重复（论文自我重复）\n")
    print(f"- 重复 n-gram 覆盖字符：{dup_chars}")
    print(f"- **内部重复率 = {rate:.2f}%**（阈值参考：>5% 需改写）\n")
    intra.sort(reverse=True)
    if intra:
        print("| 出现次数 | 片段（前 60 字） |\n|---|---|")
        for c, g in intra[:a.top]:
            print(f"| {c} | {g} |")
    print("")
    result = {"target": a.target, "nchar": total, "internal_rate": round(rate, 2)}

    if a.ref:
        files = [f for f in collect_dir(a.ref)
                 if os.path.abspath(f) != os.path.abspath(a.target)]
        ref_grams = set()
        for f in files:
            t = "".join(normalize(read_text(f)))
            ref_grams.update(ngrams(t, N))
        ref_mark = [False] * total
        for i in range(max(total - N + 1, 0)):
            if body[i:i + N] in ref_grams:
                for k in range(i, min(i + N, total)):
                    ref_mark[k] = True
        common = covered_chars(ref_mark)
        cov = common / total * 100 if total else 0
        print("## 2. 源库重合（源库 = 专利交底书/素材包/旧稿）\n")
        print(f"- 源库文件数：{len(files)}；源库 n-gram 数：{len(ref_grams)}")
        print(f"- 与源库重合成字符：{common}")
        print(f"- **源库重合率 = {cov:.2f}%**（阈值参考：>10% 建议重述）\n")
        blocks = longest_common_blocks(body, ref_grams, N, max(N, 20), a.top)
        if blocks:
            print("### 最长重合片段 Top（优先改这些）\n")
            print("| 长度 | 正文位置 | 片段 |\n|---|---|---|")
            for ln, pos, seg in blocks:
                show = seg if len(seg) <= 120 else seg[:120] + "…"
                print(f"| {ln} | {pos} | {show} |")
        else:
            print("未发现 ≥20 字的连续重合片段。")
        print("")
        result.update({"ref_files": len(files), "ref_rate": round(cov, 2),
                       "longest_blocks": [{"len": l, "pos": pos} for l, pos, _ in blocks]})

    print("---")
    print("⚠ **口径声明**：本结果为自建库 n-gram 自检（n=%d），**不是知网/CNKI 查重**，不可作为官方重复率汇报；仅用于定位需重述的片段。" % N)
    if a.json:
        ensure_parent(a.json)
        with open(a.json, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
        print(f"\n[写盘] {a.json}")


if __name__ == "__main__":
    main()
