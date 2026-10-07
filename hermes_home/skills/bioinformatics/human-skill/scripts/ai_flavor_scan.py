#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""AI 腔扫描器（中文学术文本）

用法:
    python ai_flavor_scan.py <file.md|.txt|.docx> [--out report.md] [--json out.json]

诊断四类指标并给出行号：
  1. AI 腔词表命中（必删级/慎用级）
  2. 模板句式命中（不是..而是../不仅..而且../随着..的不断发展 ...）
  3. 节奏均匀度（句长/段长方差 —— 越低越像 AI）
  4. 格式噪音（破折号、加粗、emoji、弯引号、项目符号）
只诊断，不改写。阈值参考：词表密度 > 3 次/千字需处理。
"""
import os
import re
import sys
import json
import argparse
from collections import Counter


def load(path):
    """读 md/txt 或 docx（docx 用 python-docx，失败则报错退出）。"""
    if path.lower().endswith(".docx"):
        try:
            from docx import Document
        except ImportError:
            sys.exit("需要 python-docx 读 .docx：请用装了 python-docx 的解释器运行（或用 Markdown 母本）")
        d = Document(path)
        lines = []
        for p in d.paragraphs:
            lines.append(p.text)
        for t in d.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text for c in row.cells))
        return lines
    with open(path, encoding="utf-8", errors="ignore") as f:
        return f.read().splitlines()


# ---- 词表 ----------------------------------------------------------------
MUST_DELETE = [
    "赋能", "抓手", "闭环", "加持", "深度融合", "全面提升", "极大地", "显著地",
    "里程碑式", "开创性", "颠覆性", "值得深入探讨", "具有重要意义", "重要里程碑",
    "开创了", "填补了空白", "引领了", "全新范式",
]
CAUTION = [
    "此外", "综上", "不仅", "而且", "一方面", "另一方面",
    "呈现出", "的趋势", "在.*背景下", "提供了新思路", "拓宽了",
    "奠定了坚实基础", "凸显了", "彰显了", "充分体现了",
]
TEMPLATES = {
    "不是…而是…": r"不是[^，。；]{1,25}[，,]?\s*而是",
    "不仅…而且…": r"不仅[^，。；]{1,30}[，,]?\s*(而且|还|也)",
    "随着…的不断发展": r"随着[^，。；]{0,20}(不断|持续|日益)(发展|进步|深入|推进)",
    "值得注意的是": r"值得注意的是",
    "综上所述/总而言之": r"(综上所述|总而言之|总的来说)",
    "在…背景下": r"在[^，。；]{0,20}的?背景下",
    "为…提供了新思路": r"为[^，。；]{0,25}提供了(新|全新)(思路|视角|方向)",
    "从而/进而 制造深度": r"[，,]\s*(从而|进而|以此)",
    "凸显了/彰显了": r"(凸显了|彰显了|体现了[^，。；]{0,10}的重要性)",
    "从…到…（虚假区间）": r"从[^，。；]{1,20}到[^，。；]{1,20}[，,]?\s*(从[^，。；]{1,20}到[^，。；]{1,20})",
}
NOISE = {
    "破折号 —": r"——",
    "emoji": r"[\U0001F300-\U0001FAFF\u2600-\u27BF]",
    "弯曲引号": r"[\u2018\u2019\u201C\u201D]",
    "Markdown 加粗": r"\*\*",
    "项目符号行": r"^\s*[-*+•]\s+",
}


def norm_lines(lines):
    """剥离图片/表格线/代码块，留下正文（带原始行号）。"""
    out = []
    in_code = False
    for i, L in enumerate(lines, 1):
        s = L.strip()
        if s.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            continue
        if s.startswith("!["):          # 图片
            continue
        if re.match(r"^\|.*\|$", s):    # 表格行
            continue
        if not s:
            continue
        out.append((i, L))
    return out


def sentences(text):
    """按中文句末标点切句。"""
    return [s.strip() for s in re.split(r"[。！？；!?;]", text) if len(s.strip()) >= 2]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--out", default=None)
    ap.add_argument("--json", default=None)
    a = ap.parse_args()

    raw = load(a.path)
    body = norm_lines(raw)
    fulltext = "\n".join(L for _, L in body)
    nchar = len(re.sub(r"\s", "", fulltext))
    per_k = max(nchar / 1000.0, 1e-9)
    rep = []

    rep.append(f"# AI 腔扫描报告\n\n**文件**：`{a.path}`\n**正文字符数**：{nchar}（按 {per_k:.1f} 千字折算）\n")

    # 1. 词表
    rep.append("## 1. AI 腔词表命中\n")
    rows, cnt = [], Counter()
    for i, L in body:
        for w in MUST_DELETE:
            if w in L:
                cnt[w] += L.count(w)
                rows.append(("必删", w, i, L.strip()[:80]))
        for w in CAUTION:
            if re.search(w, L):
                key = re.sub(r"[.*]", "", w) or w
                cnt[key] += 1
                rows.append(("慎用", key, i, L.strip()[:80]))
    total = sum(cnt.values())
    rep.append(f"命中 **{total}** 次，密度 **{total/per_k:.2f}/千字**（阈值 3.00）\n")
    if rows:
        rep.append("| 级别 | 词 | 行号 | 上下文 |\n|---|---|---|---|")
        for lv, w, i, ctx in rows[:120]:
            rep.append(f"| {lv} | {w} | {i} | {ctx.replace('|','/')} |")
    else:
        rep.append("无命中。")
    rep.append("")

    # 2. 模板句
    rep.append("## 2. 模板句式命中\n")
    trows = []
    for name, pat in TEMPLATES.items():
        for i, L in body:
            for m in re.finditer(pat, L):
                trows.append((name, i, m.group(0)))
    rep.append(f"命中 **{len(trows)}** 处\n")
    if trows:
        rep.append("| 句式 | 行号 | 命中片段 |\n|---|---|---|")
        for name, i, s in trows[:120]:
            rep.append(f"| {name} | {i} | {s.replace('|','/')} |")
    rep.append("")

    # 3. 节奏
    rep.append("## 3. 节奏均匀度\n")
    sents = [x for _, _L in body for x in sentences(_L)]
    slen = [len(s) for s in sents]
    paras = [len(re.sub(r"\s", "", L)) for _, L in body if len(re.sub(r"\s", "", L)) > 20]
    def stats(v):
        if not v:
            return (0, 0, 0)
        mu = sum(v) / len(v)
        sd = (sum((x - mu) ** 2 for x in v) / len(v)) ** 0.5
        return (round(mu, 1), round(sd, 1), round(sd / mu, 3) if mu else 0)
    sm, ssd, scv = stats(slen)
    pm, psd, pcv = stats(paras)
    rep.append(f"- 句数 {len(sents)}｜均长 {sm} 字｜标准差 {ssd}｜变异系数 **{scv}**")
    rep.append(f"- 段数 {len(paras)}｜均长 {pm} 字｜标准差 {psd}｜变异系数 **{pcv}**")
    rep.append(f"- 判读：句长变异系数 < 0.45 或段长变异系数 < 0.40 → 节奏过于均匀，AI 特征明显\n")

    # 4. 噪音
    rep.append("## 4. 格式噪音\n")
    nrows = []
    for name, pat in NOISE.items():
        c = len(re.findall(pat, fulltext, re.M))
        nrows.append((name, c))
    if nrows:
        rep.append("| 项 | 次数 |\n|---|---|")
        for name, c in nrows:
            rep.append(f"| {name} | {c} |")
    rep.append("")

    report = "\n".join(rep)
    print(report)
    if a.out:
        os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
        with open(a.out, "w", encoding="utf-8") as f:
            f.write(report)
        print(f"\n[写盘] {a.out}")
    if a.json:
        os.makedirs(os.path.dirname(os.path.abspath(a.json)), exist_ok=True)
        with open(a.json, "w", encoding="utf-8") as f:
            json.dump({"nchar": nchar, "word_hits": total, "template_hits": len(trows),
                       "sent_cv": scv, "para_cv": pcv, "noise": dict(nrows)},
                      f, ensure_ascii=False, indent=2)
        print(f"[写盘] {a.json}")


if __name__ == "__main__":
    main()
